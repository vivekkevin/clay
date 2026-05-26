#!/usr/bin/env python3
"""
find_addresses_gemini.py
========================
Find the complete street address for each location row in a spreadsheet.
Same input / same output as the Anthropic version — only the LLM backend changed
to a LiteLLM (OpenAI-compatible) endpoint serving Gemini Flash.

INPUT  : an .xlsx with at least two columns:
           - "Text"        -> "Company Name, Location Name, City, State, Country"
           - "Source Link" -> a URL pointing to a source document
OUTPUT : an .xlsx with columns:
           S.No | Text | Source Link | Full Address | Status | Notes

Two modes
---------
  heuristic (default) : pure-Python. Fetches each Source Link and pulls an
                        address-shaped string from JSON-LD / regex. No key needed.
  llm                 : sends the fetched page text to Gemini Flash (via LiteLLM)
                        and asks it to extract the correct street address as JSON.
                        NOTE: most LiteLLM gateways do NOT expose a server-side
                        web_search tool, so this mode extracts from the page you
                        fetched; broken/careers/login links fall back to the
                        city/state/country partial (and say so in Notes).

Config (env vars override these defaults)
-----------------------------------------
  LITELLM_BASE_URL   default https://chatapi.straive.com
                     (set to http://localhost:4000 to use a local LiteLLM proxy)
  LITELLM_MODEL      default my-gemini-model
  LITELLM_API_KEY    your gateway key  (REQUIRED for --mode llm)

Usage
-----
  pip install requests beautifulsoup4 openpyxl pandas lxml openai pypdf

  # heuristic, no key:
  python find_addresses_gemini.py input.xlsx -o output.xlsx

  # Gemini Flash via LiteLLM:
  #   PowerShell:  $env:LITELLM_API_KEY="sk-..."
  #   bash/zsh:    export LITELLM_API_KEY=sk-...
  python find_addresses_gemini.py input.xlsx -o output.xlsx --mode llm

  --limit 20         only process the first 20 rows (test run)
  --start 100        skip the first 100 rows (resume)
  --sleep 1.5        seconds between rows
  --model NAME       override LITELLM_MODEL for this run
  --base-url URL     override LITELLM_BASE_URL for this run
"""

import argparse
import json
import os
import re
import sys
import time
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ----------------------------------------------------------------------------- #
# Configuration
# ----------------------------------------------------------------------------- #

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}
REQUEST_TIMEOUT = 25

SKIP_URL_HINTS = (
    "careers", "career", "jobs", "job-", "/job/", "login", "signin", "sign-in",
    "auth", "account", "linkedin.com/jobs", "indeed.com", "glassdoor",
)

OUTPUT_COLUMNS = ["S.No", "Text", "Source Link", "Full Address", "Status", "Notes"]

# Token-budget constants
PAGE_EXCERPT_CHARS = 6_000   # address info is almost always near the top
LLM_MAX_TOKENS     = 300     # the JSON reply is tiny

# --- LiteLLM / Gemini Flash endpoint defaults (override via env or CLI) ------- #
DEFAULT_BASE_URL = os.environ.get("LITELLM_BASE_URL", "https://chatapi.straive.com")
DEFAULT_MODEL    = os.environ.get("LITELLM_MODEL", "my-gemini-model")
# API key is read from LITELLM_API_KEY at call time — never hardcode it here.


# ----------------------------------------------------------------------------- #
# Parsing the input "Text" field
# ----------------------------------------------------------------------------- #

def parse_text_field(text):
    parts = [p.strip() for p in str(text).split(",") if p.strip()]
    fields = {"company": "", "location": "", "city": "", "state": "", "country": "",
              "raw": str(text)}
    keys = ["company", "location", "city", "state", "country"]
    for i, key in enumerate(keys):
        if i < len(parts):
            fields[key] = parts[i]
    if parts:
        fields["country"] = parts[-1]
    return fields


def looks_like_skip_url(url):
    u = (url or "").lower()
    if not u or not u.startswith(("http://", "https://")):
        return True
    return any(h in u for h in SKIP_URL_HINTS)


# ----------------------------------------------------------------------------- #
# Fetching pages
# ----------------------------------------------------------------------------- #

def fetch(url):
    try:
        r = requests.get(url, headers=HEADERS, timeout=REQUEST_TIMEOUT,
                         allow_redirects=True)
        ctype = r.headers.get("Content-Type", "").lower()
        if r.status_code >= 400:
            return None, ctype, f"HTTP {r.status_code}"
        if "application/pdf" in ctype or url.lower().endswith(".pdf"):
            return _extract_pdf_text(r.content), "pdf", None
        return r.text, ctype, None
    except requests.RequestException as e:
        return None, "", str(e)


def _extract_pdf_text(content):
    try:
        import io
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        return ""


# ----------------------------------------------------------------------------- #
# Page-text cleaning — cuts token waste before sending to the LLM
# ----------------------------------------------------------------------------- #

_NOISE_RE = re.compile(
    r"(?:nav(?:igation)?|header|footer|cookie|privacy policy|terms of (?:use|service)"
    r"|all rights reserved|©|\bskip to\b)",
    re.IGNORECASE,
)

def clean_page_text(raw_text: str) -> str:
    lines = raw_text.splitlines()
    kept = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped) < 4:
            continue
        if _NOISE_RE.search(stripped):
            continue
        kept.append(stripped)

    deduped = []
    prev = None
    for ln in kept:
        if ln != prev:
            deduped.append(ln)
        prev = ln

    return " ".join(deduped)


# ----------------------------------------------------------------------------- #
# Heuristic address extraction
# ----------------------------------------------------------------------------- #

STREET_SUFFIXES = (
    r"Street|St\.?|Avenue|Ave\.?|Road|Rd\.?|Boulevard|Blvd\.?|Lane|Ln\.?|Drive|Dr\.?|"
    r"Way|Court|Ct\.?|Plaza|Square|Sq\.?|Suite|Ste\.?|Floor|Fl\.?|Highway|Hwy\.?|"
    r"Parkway|Pkwy\.?|Place|Pl\.?|Terrace|Circle|Cir\.?|Marg|Nagar|Road|Park"
)
STREET_RE = re.compile(
    r"\d{1,5}[\w\.\-,/ ]{2,60}?\b(?:%s)\b[\w\.\-,/# ]{0,80}" % STREET_SUFFIXES,
    re.IGNORECASE,
)
POSTAL_RE = re.compile(r"\b(?:\d{5}(?:-\d{4})?|[A-Z]\d[A-Z]\s?\d[A-Z]\d|\d{6})\b")


def extract_from_jsonld(soup):
    found = []
    for tag in soup.find_all("script", type="application/ld+json"):
        try:
            data = json.loads(tag.string or "")
        except Exception:
            continue
        for obj in _iter_objs(data):
            addr = obj.get("address") if isinstance(obj, dict) else None
            if isinstance(addr, dict) and addr.get("@type", "").endswith("PostalAddress"):
                pieces = [
                    addr.get("streetAddress"), addr.get("addressLocality"),
                    addr.get("addressRegion"), addr.get("postalCode"),
                    addr.get("addressCountry"),
                ]
                joined = ", ".join(_flatten(pieces))
                if joined.strip():
                    found.append(joined.strip())
    return found


def _iter_objs(data):
    if isinstance(data, list):
        for d in data:
            yield from _iter_objs(d)
    elif isinstance(data, dict):
        yield data
        for v in data.values():
            if isinstance(v, (list, dict)):
                yield from _iter_objs(v)


def _flatten(pieces):
    out = []
    for p in pieces:
        if isinstance(p, dict):
            p = p.get("name") or p.get("@id") or ""
        if p:
            out.append(str(p).strip())
    return out


def heuristic_extract(page_text, content_type):
    if not page_text:
        return None, None

    if content_type == "pdf":
        text = page_text
        soup = None
    else:
        soup = BeautifulSoup(page_text, "lxml")
        jsonld = extract_from_jsonld(soup)
        if jsonld:
            return jsonld[0], "street"
        text = soup.get_text(" ", strip=True)

    m = STREET_RE.search(text)
    if m:
        snippet = m.group(0).strip(" ,;")
        tail = text[m.end(): m.end() + 40]
        pm = POSTAL_RE.search(tail)
        if pm and pm.group(0) not in snippet:
            snippet = f"{snippet}, {pm.group(0)}"
        return snippet, "street"

    return None, None


# ----------------------------------------------------------------------------- #
# LLM-assisted extraction — Gemini Flash via LiteLLM (OpenAI-compatible API)
# ----------------------------------------------------------------------------- #

_LLM_SYSTEM = (
    "Find the exact published street address of the given company location. "
    "Rules: never invent any part of an address; copy it exactly as published. "
    "Use ONLY the page text provided — if it lacks a usable street address, set "
    "status to 'Partial' (if you can confirm city/region/postal/country) or "
    "'Not Found'. Do not guess. "
    "status: 'Verified' if a full street address is present in the page text; "
    "'Partial' if only city/region/postal/country; 'Not Found' if nothing reliable. "
    'Reply ONLY with JSON: {"address":"...","status":"Verified|Partial|Not Found","notes":"..."} '
    "In notes, mention any caveats or alternatives."
)


def _get_client(base_url):
    try:
        from openai import OpenAI
    except ImportError:
        raise SystemExit("pip install openai  (needed for --mode llm)")
    api_key = os.environ.get("LITELLM_API_KEY")
    if not api_key:
        raise SystemExit("Set LITELLM_API_KEY for --mode llm.")
    return OpenAI(base_url=base_url, api_key=api_key)


def llm_extract(fields, url, page_text, content_type, model, base_url):
    cleaned = clean_page_text(page_text or "")[:PAGE_EXCERPT_CHARS]
    if not cleaned:
        # No page text to reason over (broken/careers/login/empty PDF).
        # Skip the API call entirely and fall back to the input partial.
        return _partial_or_notfound(
            fields,
            "Source link unusable (broken/careers/login/empty); no page text to read. "
            "Needs an internet search for the full street address.")

    client = _get_client(base_url)

    user = (
        f"Company: {fields['company']}\n"
        f"Location: {fields['location']}, {fields['city']}, "
        f"{fields['state']}, {fields['country']}\n"
        f"Source: {url}\n\n"
        f"Page text ({content_type}):\n{cleaned}\n\n"
        "Return only the JSON object."
    )

    resp = client.chat.completions.create(
        model=model,
        max_tokens=LLM_MAX_TOKENS,
        temperature=0,
        messages=[
            {"role": "system", "content": _LLM_SYSTEM},
            {"role": "user", "content": user},
        ],
    )
    text_out = resp.choices[0].message.content or ""
    return _parse_llm_json(text_out)


def _parse_llm_json(text):
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    try:
        obj = json.loads(text)
    except Exception:
        m = re.search(r"\{.*\}", text, re.DOTALL)
        obj = json.loads(m.group(0)) if m else {}
    return {
        "address": str(obj.get("address", "") or "").strip(),
        "status":  str(obj.get("status",  "Not Found") or "Not Found").strip(),
        "notes":   str(obj.get("notes",   "") or "").strip(),
    }


# ----------------------------------------------------------------------------- #
# Per-row processing
# ----------------------------------------------------------------------------- #

def _partial_or_notfound(fields, lead_note):
    bits = [b for b in (fields["city"], fields["state"], fields["country"]) if b]
    if bits:
        return {"address": ", ".join(bits), "status": "Partial",
                "notes": lead_note + " Filled city/state/country from input."}
    return {"address": "", "status": "Not Found", "notes": lead_note}


def process_row_heuristic(fields, url):
    if looks_like_skip_url(url):
        return _partial_or_notfound(
            fields,
            "Source link is a careers/login/broken URL; needs internet search "
            "(research manually).")

    page_text, ctype, err = fetch(url)
    if err:
        return _partial_or_notfound(fields, f"Could not fetch source link ({err}).")

    addr, conf = heuristic_extract(page_text, ctype)
    if addr and conf == "street":
        return {"address": addr, "status": "Verified",
                "notes": "Auto-extracted from source link; please verify."}

    return _partial_or_notfound(fields, "No street address found on source page.")


def process_row_llm(fields, url, model, base_url):
    page_text, ctype = "", ""
    if not looks_like_skip_url(url):
        page_text, ctype, _err = fetch(url)
    try:
        result = llm_extract(fields, url, page_text or "", ctype, model, base_url)
    except SystemExit:
        raise
    except Exception as e:
        result = process_row_heuristic(fields, url)
        result["notes"] = f"LLM call failed ({e}); used heuristic. " + result["notes"]

    if result["status"] not in ("Verified", "Partial", "Not Found"):
        result["status"] = "Not Found" if not result["address"] else "Partial"
    if not result["address"]:
        result["status"] = "Not Found"
    return result


# ----------------------------------------------------------------------------- #
# I/O
# ----------------------------------------------------------------------------- #

def read_input(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("Input sheet is empty.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]

    def find_col(*names):
        for i, h in enumerate(header):
            if h.lower() in [n.lower() for n in names]:
                return i
        return None

    text_i = find_col("Text")
    link_i = find_col("Source Link", "Source", "Link", "URL")
    if text_i is None or link_i is None:
        raise SystemExit(
            f"Could not find required columns. Found: {header}\n"
            "Need a 'Text' column and a 'Source Link' column."
        )
    data = []
    for r in rows[1:]:
        text = r[text_i] if text_i < len(r) else ""
        link = r[link_i] if link_i < len(r) else ""
        if text is None and link is None:
            continue
        data.append((str(text or ""), str(link or "")))
    return data


def write_output(path, results):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Addresses"

    header_fill = PatternFill("solid", start_color="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    body_font   = Font(name="Arial", size=10)
    thin   = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap   = Alignment(vertical="top", wrap_text=True)

    status_fill = {
        "Verified":  PatternFill("solid", start_color="E2EFDA"),
        "Partial":   PatternFill("solid", start_color="FFF2CC"),
        "Not Found": PatternFill("solid", start_color="FCE4E4"),
    }

    sheet.append(OUTPUT_COLUMNS)
    for c in sheet[1]:
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(vertical="center", horizontal="center")
        c.border = border

    for row in results:
        sheet.append([row[c] for c in OUTPUT_COLUMNS])
        r = sheet.max_row
        for c in sheet[r]:
            c.font = body_font
            c.alignment = wrap
            c.border = border
        sf = status_fill.get(row["Status"])
        if sf:
            sheet.cell(row=r, column=5).fill = sf

    widths = {"A": 7, "B": 45, "C": 50, "D": 50, "E": 12, "F": 45}
    for col, w in widths.items():
        sheet.column_dimensions[col].width = w
    sheet.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------- #
# Main
# ----------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser(description="Find street addresses for location rows.")
    ap.add_argument("input",  help="Input .xlsx path")
    ap.add_argument("-o", "--output", default="addresses_output.xlsx")
    ap.add_argument("--mode",  choices=["heuristic", "llm"], default="heuristic")
    ap.add_argument("--model", default=DEFAULT_MODEL,
                    help=f"LiteLLM model name (default {DEFAULT_MODEL})")
    ap.add_argument("--base-url", default=DEFAULT_BASE_URL,
                    help=f"LiteLLM base URL (default {DEFAULT_BASE_URL})")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--start", type=int, default=0)
    ap.add_argument("--sleep", type=float, default=1.0)
    args = ap.parse_args()

    if args.mode == "llm" and not os.environ.get("LITELLM_API_KEY"):
        raise SystemExit("Set LITELLM_API_KEY for --mode llm.")

    data = read_input(args.input)
    if args.start:
        data = data[args.start:]
    if args.limit:
        data = data[: args.limit]

    results = []
    total = len(data)
    for idx, (text, link) in enumerate(data, start=args.start + 1):
        fields = parse_text_field(text)
        print(f"[{idx}/{args.start + total}] {fields['company'][:40]} ... ",
              end="", flush=True)
        try:
            if args.mode == "llm":
                res = process_row_llm(fields, link, args.model, args.base_url)
            else:
                res = process_row_heuristic(fields, link)
        except SystemExit:
            raise
        except Exception as e:
            res = {"address": "", "status": "Not Found", "notes": f"Row error: {e}"}

        results.append({
            "S.No":         idx,
            "Text":         text,
            "Source Link":  link,
            "Full Address": res["address"],
            "Status":       res["status"],
            "Notes":        res["notes"],
        })
        print(res["status"])
        if idx % 25 == 0:
            write_output(args.output, results)
        time.sleep(args.sleep)

    write_output(args.output, results)
    print(f"\nDone. Wrote {len(results)} rows to {args.output}")
    from collections import Counter
    print("Summary:", dict(Counter(r["Status"] for r in results)))


if __name__ == "__main__":
    main()
